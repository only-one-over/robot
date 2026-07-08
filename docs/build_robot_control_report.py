from pathlib import Path
from copy import deepcopy
from math import ceil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "robot_control_technology_research_report_formula_corrected.docx"
FIGURES = ROOT / "figures"

TITLE_CN = "六轴机械臂运动学与仿真控制"
TITLE_EN = "Kinematics and Simulation Control of a Six-Axis Robot Arm"
AUTHOR = "朱家辉"
STUDENT_ID = "202504703066"
COURSE_NAME = "机器人控制技术"

ABSTRACT_CN = (
    "针对六自由度机械臂数学模型、机器人描述与仿真控制的坐标一致性问题，构建URDF可追溯变换链，"
    "并以齐次矩阵和对偶四元数交叉验证正运动学。在统一Jacobian、初值、限位和阈值下，比较"
    "Jacobian转置、Moore-Penrose伪逆与阻尼最小二乘三种逆解。1000组样本的"
    "成功率依次为87.3%、98.7%和100.0%，平均耗时为22.29、1.68和0.62 ms；两种正解最大位置差"
    "为6.49×10^-16 m。基于ROS 2与Gazebo完成不可达目标拒绝、轨迹执行和RViz交互控制。"
    "结果表明，阻尼最小二乘在所设条件下具有最佳综合性能，所建框架支持机械臂运动学算法的"
    "可复现实验评价。"
)
KEYWORDS_CN = ["六自由度机械臂", "逆运动学", "阻尼最小二乘", "对偶四元数", "ROS 2", "Gazebo"]

ABSTRACT_EN = (
    "A traceable kinematic chain is required to keep the mathematical model, robot "
    "description, and simulated control of a six-degree-of-freedom serial manipulator "
    "consistent. This study constructs a URDF-aligned local transformation chain and "
    "cross-validates forward kinematics using homogeneous matrices and dual quaternions. "
    "With the same finite-difference Jacobian, initial perturbations, joint limits, and "
    "convergence thresholds, three inverse-kinematics updates are compared: Jacobian "
    "transpose, Moore-Penrose pseudoinverse, and damped least squares. Across 1,000 random "
    "samples, their success rates are 87.3%, 98.7%, and 100.0%, and their mean runtimes are "
    "22.29, 1.68, and 0.62 ms, respectively. The maximum positional discrepancy between "
    "matrix and dual-quaternion forward kinematics is 6.49 × 10^-16 m. ROS 2 Jazzy, "
    "ros2_control, and Gazebo are used to verify unreachable-target rejection, trajectory "
    "execution, and RViz interactive control. Under the tested local-perturbation and "
    "tolerance conditions, damped least squares provides the strongest overall balance. "
    "The resulting framework supports reproducible evaluation of serial-manipulator "
    "kinematics algorithms."
)
KEYWORDS_EN = [
    "six-degree-of-freedom manipulator",
    "inverse kinematics",
    "damped least squares",
    "dual quaternion",
    "ROS 2",
    "Gazebo",
]

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = "E7E7E7"
MID_GRAY = "B7B7B7"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5,
                 bold=None, italic=None, color=BLACK):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def set_table_geometry(table, widths_mm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_twips = int(sum(widths_mm) / 25.4 * 1440)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_mm in widths_mm:
        width_twips = int(width_mm / 25.4 * 1440)
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_twips))
        grid.append(grid_col)

    for row in table.rows:
        for index, (cell, width_mm) in enumerate(zip(row.cells, widths_mm)):
            width_twips = int(width_mm / 25.4 * 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Mm(width_mm)
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_cell_text(cell, text, bold=False, size=8.5,
                        align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines


def add_body(doc, text, first_indent=True, size=10.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.5
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Pt(size * 2)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    set_paragraph_keep(paragraph)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(
        run,
        east_asia="黑体",
        latin="Times New Roman",
        size=14 if level == 1 else 12,
        bold=True,
    )
    return paragraph


def add_caption(doc, text, keep_next=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=9)
    return paragraph


def add_figure(doc, filename, caption, width_mm=150):
    path = FIGURES / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Mm(width_mm))
    add_caption(doc, caption)


def add_standard_table(doc, headers, rows, widths_mm, caption):
    add_caption(doc, caption, keep_next=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, widths_mm)
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        set_table_cell_text(cell, header, bold=True, size=8.5)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), LIGHT_GRAY)
        cell._tc.get_or_add_tcPr().append(shading)
    for row_values in rows:
        row = table.add_row()
        for index, (cell, value) in enumerate(zip(row.cells, row_values)):
            if isinstance(value, dict) and "equation" in value:
                set_table_cell_equation(cell, value["equation"])
                continue
            align = WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            set_table_cell_text(cell, value, size=8.2, align=align)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    return table


def math_run(text, size_half_points=None):
    run = OxmlElement("m:r")
    r_pr = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    r_pr.append(style)
    run.append(r_pr)
    if size_half_points is not None:
        word_r_pr = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(size_half_points))
        word_r_pr.append(size)
        size_complex = OxmlElement("w:szCs")
        size_complex.set(qn("w:val"), str(size_half_points))
        word_r_pr.append(size_complex)
        run.append(word_r_pr)
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    return run


def build_equation(text, size_half_points=None):
    equation = OxmlElement("m:oMath")
    equation.append(math_run(text, size_half_points=size_half_points))
    return equation


def set_table_cell_equation(cell, text):
    cell.text = ""
    set_cell_margins(cell, top=70, start=70, bottom=70, end=70)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph._p.append(build_equation(text, size_half_points=17))


def add_equation(doc, number, text, explanation):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [15, 130, 15])
    remove_table_borders(table)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
    middle = table.cell(0, 1).paragraphs[0]
    middle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    middle.paragraph_format.space_before = Pt(2)
    middle.paragraph_format.space_after = Pt(2)
    middle.paragraph_format.keep_together = True
    middle._p.append(build_equation(text))
    right = table.cell(0, 2).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_before = Pt(2)
    right.paragraph_format.space_after = Pt(2)
    run = right.add_run(f"({number})")
    set_run_font(run, size=10.5)
    add_body(doc, f"式（{number}）中，{explanation}", first_indent=True)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, value, end))
    set_run_font(run, size=9)


def set_section_geometry(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(12)


def add_cover(doc):
    section = doc.sections[0]
    set_section_geometry(section)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("武汉科技大学")
    set_run_font(run, east_asia="黑体", size=24, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run("研究生课程研究报告")
    set_run_font(run, east_asia="黑体", size=20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(38)
    run = p.add_run(f"《{COURSE_NAME}》")
    set_run_font(run, east_asia="黑体", size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run(TITLE_CN)
    set_run_font(run, east_asia="黑体", size=22, bold=True)

    info = doc.add_table(rows=6, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(info, [42, 80])
    remove_table_borders(info)
    fields = [
        ("学生姓名", AUTHOR),
        ("学号", STUDENT_ID),
        ("学院", "________________________"),
        ("专业", "________________________"),
        ("任课教师", "________________________"),
        ("完成日期", "2026年6月"),
    ]
    for row, (label, value) in zip(info.rows, fields):
        set_table_cell_text(row.cells[0], label, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_table_cell_text(row.cells[1], value, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(
            row.cells[1],
            bottom={"val": "single", "sz": "6", "color": "000000"},
        )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("评阅记录")
    set_run_font(run, east_asia="黑体", size=11, bold=True)

    review = doc.add_table(rows=2, cols=5)
    review.style = "Table Grid"
    set_table_geometry(review, [32, 32, 32, 32, 32])
    labels = ["选题与资料\n（10分）", "正文内容\n（60分）", "报告表述\n（10分）", "创新性\n（20分）", "总分\n（100分）"]
    for cell, label in zip(review.rows[0].cells, labels):
        set_table_cell_text(cell, label, bold=True, size=8.5)
    for cell in review.rows[1].cells:
        set_table_cell_text(cell, "        ", size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("评语：")
    set_run_font(run, east_asia="黑体", size=10, bold=True)
    for _ in range(3):
        line = doc.add_paragraph("________________________________________________________________")
        line.paragraph_format.space_before = Pt(0)
        line.paragraph_format.space_after = Pt(0)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(line.runs[0], size=9)


def configure_body_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), "1")


def add_article_front(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(TITLE_CN)
    set_run_font(run, east_asia="黑体", size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(TITLE_EN)
    set_run_font(run, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(AUTHOR)
    set_run_font(run, size=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("（武汉科技大学，湖北 武汉 430065）")
    set_run_font(run, size=9)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    label = p.add_run("摘要：")
    set_run_font(label, east_asia="黑体", size=10.5, bold=True)
    run = p.add_run(ABSTRACT_CN)
    set_run_font(run, size=10.5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.35

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    label = p.add_run("关键词：")
    set_run_font(label, east_asia="黑体", size=10.5, bold=True)
    run = p.add_run("；".join(KEYWORDS_CN))
    set_run_font(run, size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    label = p.add_run("Abstract: ")
    set_run_font(label, size=10, bold=True)
    run = p.add_run(ABSTRACT_EN)
    set_run_font(run, size=10)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    label = p.add_run("Key words: ")
    set_run_font(label, size=10, bold=True)
    run = p.add_run("; ".join(KEYWORDS_EN))
    set_run_font(run, size=10)


def add_main_text(doc):
    add_heading(doc, "1 引言", 1)
    add_body(
        doc,
        "串联机械臂通过关节变量改变末端执行器在任务空间中的位置和姿态，其运动学模型是轨迹规划、"
        "伺服控制与碰撞检测的共同基础。Denavit-Hartenberg（DH）参数化为相邻连杆坐标关系提供了"
        "经典表达[1]，现代机器人学进一步从矩阵群与旋量角度统一了刚体运动描述[2-3]。然而，数学"
        "模型正确并不等同于软件系统正确：若运动学坐标系、URDF父子连杆关系和控制器关节方向不一致，"
        "数值逆解即使在自身公式中收敛，也会在仿真执行阶段产生系统性偏差。"
    )
    add_body(
        doc,
        "逆运动学的核心困难来自非线性、多解性、关节限位与奇异位形。Whitney提出基于Jacobian的"
        "分辨率运动速率控制[4]；Nakamura和Hanafusa以及Wampler分别发展了奇异鲁棒逆解与阻尼最小"
        "二乘方法[5-6]。后续研究从任务优先级、选择性阻尼和Levenberg-Marquardt正则化等方向提高"
        "数值稳定性[7-9]。Jacobian转置、伪逆与阻尼最小二乘均具有清晰的线性化解释，但其收敛率、"
        "计算代价和奇异邻域行为依赖于具体机构、初值分布与停止阈值，因此不能脱离统一实验条件比较。"
    )
    add_body(
        doc,
        "近年研究继续从高阶迭代、关节优先级和求解加速等方向改进通用逆解[12-15]，并将对偶四元数"
        "用于刚体位姿、运动控制和六轴机械臂解析建模[11,16-19]。ROS 2为机器人软件提供模块化通信"
        "与执行架构[20]，国内研究也围绕机械臂几何逆解、多项式轨迹和采样式路径规划开展了系统仿真"
        "分析[21-22]。现有成果分别回答了算法、表示或平台问题，但在同一URDF模型上同时进行矩阵—"
        "对偶四元数正解互验、三种Jacobian逆解公平比较和Gazebo闭环执行的研究仍具有方法整合价值。"
    )
    add_body(
        doc,
        "本文不提出新的逆运动学更新律，而是构建一套可复现的系统级评价方法。主要工作包括："
        "（1）建立与URDF joint origin和axis逐项对应的等效局部变换链；（2）以齐次矩阵和对偶四元数"
        "实现相互独立的正运动学校验；（3）在完全一致的随机样本、初值、关节限位和收敛阈值下比较"
        "Jacobian转置、Moore-Penrose伪逆和阻尼最小二乘三种逆解；（4）通过ROS 2、ros2_control、"
        "Gazebo与RViz交互标记验证从任务空间目标到关节轨迹执行的闭环。"
    )

    add_heading(doc, "2 系统建模与总体架构", 1)
    add_heading(doc, "2.1 机械臂本体与坐标语义", 2)
    add_body(
        doc,
        "研究对象为六个转动关节串联组成的机械臂，关节命名为joint_1至joint_6，基坐标系为"
        "base_link，工具坐标系为tool0。为降低网格模型对仿真速度的影响，连杆视觉与碰撞几何采用"
        "圆柱体和长方体近似；运动学计算仅依赖关节原点、旋转轴和固定工具变换，因此几何外观简化"
        "不会改变本文比较的运动学映射。"
    )
    rows = [
        ("d1", "0.0985 m", "joint_1沿局部+Z方向的基座偏置"),
        ("a2", "-0.408 m", "joint_2至joint_3的主臂局部X向长度"),
        ("a3", "-0.376 m", "joint_3至joint_4的主臂局部X向长度"),
        ("d4", "0.1215 m", "腕部局部Y向偏置"),
        ("d5", "0.1025 m", "joint_5至joint_6的局部Z向偏置"),
        ("d6", "0.0940 m", "末端工具长度"),
    ]
    add_standard_table(doc, ["参数", "数值", "几何含义"], rows, [28, 35, 97], "表 1  六轴机械臂主要几何参数")
    add_body(
        doc,
        "标准DH与改进DH（MDH）的差异在于坐标系附着规则和基本变换顺序。标准DH适合用四参数表"
        "压缩表达连杆链；MDH更强调前一连杆坐标系与当前关节之间的局部关系。本研究采用“URDF等效"
        "局部变换链”而不将其机械地称为标准MDH表：每一平移和转动均可追溯到Xacro中的joint origin"
        "与axis，其组织思想与MDH局部关节语义一致，但表达形式直接服从URDF父子连杆定义。"
    )
    add_equation(
        doc, 1,
        "Tᵢ(qᵢ) = Trans(pᵢ) · Rot(aᵢ,qᵢ),  i = 1,…,6",
        "Tᵢ为第i个关节的局部齐次变换，pᵢ为joint origin给出的平移向量，aᵢ为单位旋转轴，qᵢ为关节角。"
    )
    add_equation(
        doc, 2,
        "{}^base_link T_tool0(q) = T₁(q₁) · T₂(q₂) · T₃(q₃) · T₄(q₄) · T₅(q₅) · T₆(q₆) · {}^6 T_tool0",
        "式左侧表示末端tool0相对于base_link的位姿，q为六维关节向量；式末项表示link_6至tool0的固定变换。"
    )

    add_heading(doc, "2.2 ROS 2与仿真闭环", 2)
    add_body(
        doc,
        "系统由robot_arm_description、robot_arm_kinematics和robot_arm_bringup三个功能包组成。"
        "description包维护Xacro/URDF、惯性参数与RViz配置；kinematics包实现正逆运动学、轨迹插值、"
        "交互标记和基准测试；bringup包配置Gazebo世界、控制器管理器和joint_trajectory_controller。"
        "该分层将几何描述、算法实现与运行编排分离，使同一运动学库可以同时服务命令行实验和仿真节点。"
    )
    add_figure(
        doc,
        "fig1_system_architecture.png",
        "图 1  机械臂运动学与仿真控制总体架构",
        width_mm=150,
    )
    add_body(
        doc,
        "任务空间目标由/target_pose输入，kinematics_node根据/joint_states选择当前关节构型作为种子，"
        "逆解成功后生成五次关节轨迹，并通过FollowJointTrajectory动作接口发送至控制器。执行结果"
        "由/joint_states反馈，运动学节点重新计算/fk_pose并发布/ik_status。不可达目标在逆解层被拒绝，"
        "不会进入有效轨迹发送路径。"
    )

    add_heading(doc, "3 正运动学与位姿双表示", 1)
    add_heading(doc, "3.1 齐次矩阵正运动学", 2)
    add_body(
        doc,
        "齐次矩阵以旋转矩阵和位置向量组成SE(3)元素，能够直接呈现局部坐标变换的连乘关系。"
        "实现中forward(q)严格按照式（2）的顺序计算，负的a2和a3表示主臂沿局部负X方向展开。"
        "同一符号约定同时写入URDF和C++常量，避免通过经验性符号修正使仿真与数学模型“表面一致”。"
    )
    add_heading(doc, "3.2 对偶四元数正运动学", 2)
    add_body(
        doc,
        "对偶四元数以一个单位旋转四元数和一个对偶部统一表示旋转与平移。与4×4齐次矩阵相比，"
        "其参数更紧凑，并与螺旋位移具有直接联系[11,16-19,23]。本文将每个局部齐次变换独立转换为"
        "单位对偶四元数后再连乘，使两条正运动学路径共享几何参数而不共享矩阵乘法实现。"
    )
    add_equation(
        doc, 3,
        "Q̂ = qᵣ + εq_d,  q_d = ½p* ⊗ qᵣ,  ε² = 0",
        "Q̂为单位对偶四元数，qᵣ为单位旋转四元数，q_d为对偶部，p*为由平移向量构成的纯四元数，⊗表示四元数乘法。"
    )
    add_body(
        doc,
        "矩阵路径与对偶四元数路径的末端位置差和相对旋转角构成正解等价性指标。该交叉验证不能证明"
        "URDF本身的几何参数绝对正确，但能够高灵敏度地发现变换顺序、旋转轴符号和工具偏置在两种"
        "实现之间的不一致。"
    )
    add_figure(
        doc,
        "fig2_kinematics_flow.png",
        "图 2  齐次矩阵、对偶四元数与逆运动学误差的计算关系",
        width_mm=150,
    )

    add_heading(doc, "4 三种数值逆运动学算法", 1)
    add_heading(doc, "4.1 位姿误差与有限差分Jacobian", 2)
    add_body(
        doc,
        "三种算法共享完全相同的任务空间误差。位置误差取目标与当前末端位置之差，姿态误差采用"
        "当前旋转矩阵与目标旋转矩阵对应列向量叉乘的半和。收敛条件为位置误差小于5 mm且姿态"
        "误差小于3°；最大迭代次数为250，单关节单步增量限制为0.18 rad，更新后执行关节限位投影。"
        "本文对误差函数直接作有限差分，并记误差Jacobian为J_e=∂e/∂q。它与教材中常用的任务空间"
        "Jacobian J_x=∂x/∂q近似互为相反数，因此后续更新律保留负号。"
    )
    add_equation(
        doc, 4,
        "e(q) = [e_pᵀ  e_Rᵀ]ᵀ,  e_p = p* − p(q)",
        "e为六维位姿误差，e_p为三维位置误差，e_R为三维姿态误差，p*和p(q)分别为目标位置与当前末端位置。"
    )
    add_equation(
        doc, 5,
        "e_R = ½(r_c1 × r_t1 + r_c2 × r_t2 + r_c3 × r_t3)",
        "r_cj和r_tj分别为当前旋转矩阵与目标旋转矩阵的第j列，×表示向量叉乘。"
    )
    add_equation(
        doc, 6,
        "J_e(:,i) ≈ [e(q + δᵢ eᵢ) − e(q)] / δᵢ,  0 < |δᵢ| ≤ h,  h = 10^-5",
        "J_e(:,i)为误差Jacobian的第i列，eᵢ为第i个标准基向量。δᵢ的符号选择应保证扰动后"
        "关节仍处于限位内；接近限位时分母采用不超过h的实际扰动量。"
    )

    add_heading(doc, "4.2 Jacobian转置法", 2)
    add_body(
        doc,
        "Jacobian转置法沿误差平方范数的负梯度方向更新关节角，不需求矩阵求逆。为降低固定增益对"
        "实验结论的影响，步长采用当前线性化模型上的最速下降尺度。该方法单次更新计算简单，但当"
        "Jacobian各方向尺度差异显著时收敛速度降低。"
    )
    add_equation(
        doc, 7,
        "g = J_eᵀe,  α = (gᵀg) / [(J_eg)ᵀ(J_eg)],  Δq_JT = −αg",
        "g为目标函数梯度，Δq_JT为Jacobian转置法关节增量，α为线性化最速下降步长。"
    )

    add_heading(doc, "4.3 Moore-Penrose伪逆法", 2)
    add_body(
        doc,
        "伪逆法在局部线性模型中给出最小二乘意义下的最小范数关节增量。实现采用JacobiSVD分解"
        "6×6 Jacobian，并将小于最大奇异值10^-5倍的奇异值截断。奇异值分解提高了秩判定能力，"
        "但每次迭代的计算代价高于对称正定线性方程求解，奇异值阈值也会影响解的连续性。"
    )
    add_equation(
        doc, 8,
        "J_e = UΣVᵀ,  Δq_PINV = −J_e⁺e = −VΣ⁺Uᵀe",
        "U和V为奇异向量矩阵，Σ为奇异值对角矩阵，Σ⁺对保留奇异值取倒数，J_e⁺为Moore-Penrose伪逆。"
    )

    add_heading(doc, "4.4 阻尼最小二乘法", 2)
    add_body(
        doc,
        "阻尼最小二乘法在任务误差与关节增量之间引入Tikhonov正则项，使Jacobian接近奇异时的"
        "关节增量保持有界[5-9]。本文使用固定阻尼系数λ=0.08，以便三种方法在统一参数下比较。"
        "固定阻尼会引入偏差，因此实验同时报告成功率、残差、迭代次数和耗时，而不以单一误差指标"
        "判定优劣。"
    )
    add_equation(
        doc, 9,
        "Δq_DLS = −J_eᵀ(J_eJ_eᵀ + λ²I)^−1e,  λ = 0.08",
        "Δq_DLS为阻尼最小二乘关节增量，λ为阻尼系数，I为6阶单位矩阵。"
    )

    comparison_rows = [
        (
            "Jacobian转置",
            {"equation": "Δq_JT = −αJ_eᵀe"},
            "无需分解；单步简单",
            "收敛较慢；受尺度影响",
        ),
        (
            "Moore-Penrose伪逆",
            {"equation": "Δq_PINV = −J_e⁺e = −VΣ⁺Uᵀe"},
            "局部最小范数；精度高",
            "SVD代价较高；阈值敏感",
        ),
        (
            "阻尼最小二乘",
            {"equation": "Δq_DLS = −J_eᵀ(J_eJ_eᵀ + λ²I)⁻¹e"},
            "奇异邻域有界；稳定",
            "阻尼引入偏差；需选λ",
        ),
    ]
    add_standard_table(
        doc,
        ["算法", "更新律", "主要优势", "主要限制"],
        comparison_rows,
        [34, 60, 33, 33],
        "表 2  三种Jacobian数值逆运动学方法的理论比较",
    )

    add_heading(doc, "5 轨迹生成与程序实现", 1)
    add_heading(doc, "5.1 五次关节轨迹", 2)
    add_body(
        doc,
        "逆解获得目标关节角后，采用五次多项式时间标度在起点与终点之间插值。该时间标度在两端"
        "同时满足速度和加速度为零，避免阶跃关节指令直接进入轨迹控制器。轨迹由50个离散点组成，"
        "通过FollowJointTrajectory发送；本文评价的是运动学与轨迹接口闭环，不据此推断真实执行器"
        "的动力学跟踪性能。"
    )
    add_equation(
        doc, 10,
        "q(s) = q₀ + (10s³ − 15s⁴ + 6s⁵)(q_f − q₀),  s ∈ [0,1]",
        "q₀和q_f分别为起始与目标关节向量，s为归一化时间，括号内为五次平滑时间标度函数。"
    )
    add_heading(doc, "5.2 软件接口与失败门控", 2)
    add_body(
        doc,
        "运动学库以C++和Eigen实现，三种算法通过IkMethod枚举共享同一误差、Jacobian和限位逻辑。"
        "部署节点默认使用多初值DLS，以当前关节角、零位和两组肘部构型作为候选种子，并选择收敛"
        "且综合残差较小的解；算法比较实验则关闭多初值，仅使用同一扰动种子，以隔离更新律影响。"
        "这种区分避免将部署鲁棒性策略混入算法本体比较。"
    )
    add_body(
        doc,
        "当逆解不满足双阈值时，节点发布converged=false及位置、姿态残差，不构造有效轨迹目标。"
        "交互模式使用六自由度Interactive Marker；鼠标释放事件触发PoseStamped消息，后续仍复用"
        "同一逆解、轨迹和控制器路径，因此交互输入没有绕开被评价的算法链。"
    )
    add_figure(
        doc,
        "fig3_validation_loop.png",
        "图 3  目标位姿、数值逆解、轨迹执行与状态反馈的验证闭环",
        width_mm=150,
    )

    add_heading(doc, "6 仿真实验与结果", 1)
    add_heading(doc, "6.1 实验环境与评价协议", 2)
    add_body(
        doc,
        "实验在WSL2 Ubuntu 24.04.4 LTS、ROS 2 Jazzy和Gazebo Sim环境中执行。C++工程通过"
        "colcon build --symlink-install编译，控制器采用joint_state_broadcaster和"
        "joint_trajectory_controller。随机基准固定种子20250624，共生成1000组关节向量；目标"
        "由正运动学产生，初值在真实关节角附近施加各维独立、范围为±0.25 rad的均匀扰动。"
    )
    environment_rows = [
        ("操作系统", "WSL2 Ubuntu 24.04.4 LTS"),
        ("机器人中间件", "ROS 2 Jazzy"),
        ("仿真与控制", "Gazebo Sim；ros_gz；ros2_control"),
        ("样本设置", "1000组；随机种子20250624；初值半径0.25 rad"),
        ("收敛阈值", "位置<0.005 m；姿态<3°；最大250次迭代"),
        ("统计口径", "成功率按全部样本；误差和迭代按成功样本；耗时按全部样本"),
    ]
    add_standard_table(
        doc,
        ["项目", "设置"],
        environment_rows,
        [42, 118],
        "表 3  软件环境与实验参数",
    )

    add_heading(doc, "6.2 正运动学双表示一致性", 2)
    add_body(
        doc,
        "1000组随机关节样本中，矩阵FK与对偶四元数FK的位置误差均值为1.95×10^-16 m、最大值"
        "为6.49×10^-16 m；姿态误差均值为1.60×10^-9 rad、最大值为2.98×10^-8 rad。该数量级"
        "接近双精度浮点舍入误差，说明两条独立表示路径对同一局部变换链给出一致结果。单例关节向量"
        "[0, π/2, 0, π/2, 0, 0]的两种正解位置差为7.12×10^-17 m，姿态差为0 rad。"
    )

    add_heading(doc, "6.3 三种逆运动学算法比较", 2)
    result_rows = [
        ("Jacobian转置", "873/1000", "4.883/5.000", "0.00206/0.05148", "85.66", "22.29"),
        ("伪逆", "987/1000", "1.266/4.992", "0.00359/0.05004", "2.23", "1.68"),
        ("阻尼最小二乘", "1000/1000", "2.553/4.991", "0.00249/0.04916", "2.17", "0.62"),
    ]
    add_standard_table(
        doc,
        ["算法", "成功数", "位置误差\n均值/最大(mm)", "姿态误差\n均值/最大(rad)", "平均迭代", "平均耗时(ms)"],
        result_rows,
        [29, 24, 33, 33, 20, 21],
        "表 4  1000组随机样本的三种逆运动学算法实测结果",
    )
    add_body(
        doc,
        "DLS成功率比伪逆高1.3个百分点，比Jacobian转置高12.7个百分点。DLS平均耗时最低，原因"
        "不是其单次运算最简单，而是其平均仅需2.17次迭代；Jacobian转置虽避免矩阵分解，但平均"
        "迭代85.66次，累计耗时最高。伪逆在成功样本上的平均位置误差最低，为1.266 mm，表明DLS"
        "在本实验中的优势是可靠性与总体计算成本的平衡，而非所有误差指标均占优。"
    )
    add_figure(
        doc,
        "fig5_ik_algorithm_comparison.png",
        "图 4  相同1000样本条件下三种逆运动学算法的成功率、耗时、迭代次数与位置误差。"
        "误差和迭代仅统计收敛样本，耗时统计全部样本。",
        width_mm=160,
    )

    add_heading(doc, "6.4 不可达目标与Gazebo闭环", 2)
    add_body(
        doc,
        "将目标位置设置为(5,0,5) m后，三种算法均返回converged=false。Jacobian转置、伪逆和DLS"
        "的最终位置残差分别为6.110、6.210和6.214 m，均未误判为可达目标。该实验只证明当前停止"
        "规则能够拒绝该超工作空间样例，不等同于对任意不可达位姿的完备判定。"
    )
    add_body(
        doc,
        "Gazebo闭环实验启动两个控制器并依次执行3个预设目标，运行日志记录到action goal accepted"
        "和Goal reached, success，/joint_states、/fk_pose与/ik_status均持续更新。RViz交互实验中，"
        "/target_marker/get_interactive_markers服务可用，MOUSE_UP反馈触发/target_pose发布；实测"
        "目标取前中部常用工作位姿q=[-0.35,-1.00,1.55,-0.50,-0.35,-0.20] rad，"
        "由FK得到x=-0.545 m、y=-0.024 m、z=0.145 m；该位姿远离关节限位且不处于腕部奇异状态，"
        "仿真机械臂通过同一DLS与轨迹链响应。"
    )
    add_figure(
        doc,
        "fig4_interactive_runtime.png",
        "图 5  RViz交互目标、ROS接口状态与Gazebo控制链的运行证据",
        width_mm=150,
    )
    add_body(
        doc,
        "图6给出了实际联调截图。左侧RViz以world为固定坐标系显示RobotModel、交互标记和目标坐标轴，"
        "右侧Gazebo显示同一机械臂模型在仿真场景中的姿态，终端日志同步输出trajectory goal accepted"
        "和target pose等运行信息。该图补充了图5的流程化证据，说明交互输入、逆解节点、轨迹控制器、"
        "RViz可视化和Gazebo物理仿真之间能够形成可观察的一致闭环。"
    )
    add_figure(
        doc,
        "fig6_rviz_gazebo_runtime_screenshot.png",
        "图 6  RViz交互控制、Gazebo仿真与终端日志的同步运行截图",
        width_mm=160,
    )

    add_heading(doc, "7 讨论", 1)
    add_body(
        doc,
        "实验揭示了三种更新律在统一条件下的互补性。Jacobian转置法每次迭代只需矩阵乘法，但"
        "其梯度方向未补偿Jacobian的各向异性，导致大量样本在250次上限前无法进入5 mm阈值。"
        "伪逆法利用SVD获得局部最小范数增量，因而成功样本位置误差最低；其失败主要与局部线性化、"
        "关节限位投影和奇异值截断共同作用有关。DLS通过固定阻尼限制奇异方向增益，在本机构与"
        "初值半径下获得最高成功率和最低累计耗时。"
    )
    add_body(
        doc,
        "DLS在本批1000组样本中全部收敛，但不能被解释为全局收敛保证。目标由FK生成，因此均为"
        "理论可达位姿；有限差分误差、局部极值、关节限位和固定步长上限仍可能在其他初值下限制求解。"
        "进一步地，三种算法采用同一固定参数而未对每种方法分别调优，这提高了比较可解释性，但"
        "也意味着结果不代表各算法在最优超参数下的性能上界。"
    )
    add_body(
        doc,
        "本文的系统性贡献在于把数学表示、算法比较、软件接口和仿真反馈放入同一证据链。"
        "URDF等效局部变换链保证每个局部变换可追溯，矩阵—对偶四元数互验降低了单实现自洽而"
        "共同错误的风险，失败门控则防止不可达目标进入控制器。该方法适用于教学与研究中的算法"
        "复现实验，但模型采用简化惯性和理想关节，不能外推为真实机械臂的载荷、摩擦、柔性或"
        "实时控制性能。"
    )

    add_heading(doc, "8 结论", 1)
    add_body(
        doc,
        "本文完成了六自由度机械臂从URDF可追溯建模、正运动学双表示、三种数值逆运动学比较到"
        "ROS 2/Gazebo轨迹执行的系统研究。矩阵与对偶四元数正解在浮点精度范围内一致；1000组"
        "随机样本中，Jacobian转置、伪逆和DLS成功率分别为87.3%、98.7%和100.0%，DLS平均耗时"
        "为0.62 ms，伪逆成功样本平均位置误差最低。不可达样例被三种方法拒绝，预设目标和RViz"
        "交互目标均能通过同一控制链驱动仿真模型。结果支持在当前机构、局部扰动初值和收敛阈值"
        "下选择DLS作为部署主方法，同时保留伪逆作为高精度局部求解的对照。"
    )

    add_heading(doc, "9 展望", 1)
    add_body(
        doc,
        "后续研究应围绕三类边界展开。第一，将有限差分Jacobian替换为解析或自动微分Jacobian，"
        "并在相同硬件上分离单次迭代成本与收敛次数；第二，引入基于最小奇异值或误差范数的自适应"
        "阻尼，开展参数消融和奇异位形分层测试；第三，在真实机械臂或高保真动力学模型中加入摩擦、"
        "载荷、时延和传感噪声，比较运动学解算误差与控制跟踪误差。对偶四元数还可进一步进入任务"
        "空间误差与插值定义，从而检验双表示框架在轨迹连续性方面的作用。"
    )
    add_body(
        doc,
        "结合同类研究，论文还可继续增加四类对比。其一，用Elementary Transform Sequence或解析Jacobian"
        "建立有限差分Jacobian的精度与耗时基线[25]；其二，比较伪逆、阻尼伪逆和混合广义逆在单位尺度、"
        "病态Jacobian和关节限位附近的稳定性[26]；其三，将第6轴腕部附近的奇异位形单独抽样，引入安全"
        "Jacobian投影或选择阻尼策略，以评价临近奇异点时的速度放大风险[27]；其四，结合已有6自由度"
        "机械臂对偶四元数逆解研究，为本文数值法增加解析或半解析基线，从两条路线解释误差来源[17]。"
    )


def add_references(doc):
    references = [
        ("DENAVIT J, HARTENBERG R S. A kinematic notation for lower-pair mechanisms based on matrices[J]. Journal of Applied Mechanics, 1955, 22(2): 215-221. DOI: 10.1115/1.4011045.", 1955, True),
        ("CRAIG J J. Introduction to Robotics: Mechanics and Control[M]. 3rd ed. Upper Saddle River: Pearson Prentice Hall, 2005.", 2005, True),
        ("LYNCH K M, PARK F C. Modern Robotics: Mechanics, Planning, and Control[M]. Cambridge: Cambridge University Press, 2017.", 2017, True),
        ("WHITNEY D E. Resolved motion rate control of manipulators and human prostheses[J]. IEEE Transactions on Man-Machine Systems, 1969, 10(2): 47-53. DOI: 10.1109/TMMS.1969.299896.", 1969, True),
        ("NAKAMURA Y, HANAFUSA H. Inverse kinematic solutions with singularity robustness for robot manipulator control[J]. Journal of Dynamic Systems, Measurement, and Control, 1986, 108(3): 163-171. DOI: 10.1115/1.3143764.", 1986, True),
        ("WAMPLER C W. Manipulator inverse kinematic solutions based on vector formulations and damped least-squares methods[J]. IEEE Transactions on Systems, Man, and Cybernetics, 1986, 16(1): 93-101. DOI: 10.1109/TSMC.1986.289285.", 1986, True),
        ("CHIAVERINI S. Singularity-robust task-priority redundancy resolution for real-time kinematic control of robot manipulators[J]. IEEE Transactions on Robotics and Automation, 1997, 13(3): 398-410. DOI: 10.1109/70.585902.", 1997, True),
        ("BUSS S R, KIM J S. Selectively damped least squares for inverse kinematics[J]. Journal of Graphics Tools, 2005, 10(3): 37-49. DOI: 10.1080/2151237X.2005.10129202.", 2005, True),
        ("SUGIHARA T. Solvability-unconcerned inverse kinematics by the Levenberg-Marquardt method[J]. IEEE Transactions on Robotics, 2011, 27(5): 984-991. DOI: 10.1109/TRO.2011.2148230.", 2011, True),
        ("DULEBA I, OPALKA M. A comparison of Jacobian-based methods of inverse kinematics for serial robot manipulators[J]. International Journal of Applied Mathematics and Computer Science, 2013, 23(2): 373-382. DOI: 10.2478/amcs-2013-0028.", 2013, True),
        ("DANTAM N T. Robust and efficient forward, differential, and inverse kinematics using dual quaternions[J]. The International Journal of Robotics Research, 2021, 40(10-11): 1087-1105. DOI: 10.1177/0278364920931948.", 2021, True),
        ("LLOYD S, IRANI R A, AHMADI M. Fast and robust inverse kinematics of serial robots using Halley's method[J]. IEEE Transactions on Robotics, 2022, 38(5): 2768-2780. DOI: 10.1109/TRO.2022.3162954.", 2022, True),
        ("JESUS R C O, MOLINA L, CARVALHO E A N, et al. Singularity-free inverse kinematics with joint prioritization for manipulators[J]. Journal of Control, Automation and Electrical Systems, 2022, 33(3): 1022-1031. DOI: 10.1007/s40313-021-00860-4.", 2022, True),
        ("BODO G, DI BELLO P, TESSARI F, et al. Comparative analysis of inverse kinematics methodologies to improve the controllability of rehabilitative robotic devices[C]//2022 International Conference on Rehabilitation Robotics. Piscataway: IEEE, 2022: 1-6. DOI: 10.1109/ICORR55369.2022.9896579.", 2022, True),
        ("XIE S, SUN L, WANG Z, et al. A speedup method for solving the inverse kinematics problem of robotic manipulators[J]. International Journal of Advanced Robotic Systems, 2022, 19(3): 1-10. DOI: 10.1177/17298806221104602.", 2022, True),
        ("ABAUNZA H, CHANDRA R, OZGUR E, et al. Kinematic screws and dual quaternion based motion controllers[J]. Control Engineering Practice, 2022, 128: 105325. DOI: 10.1016/j.conengprac.2022.105325.", 2022, True),
        ("AHMED A, YU M, CHEN F. Inverse kinematic solution of 6-DOF robot-arm based on dual quaternions and axis invariant methods[J]. Arabian Journal for Science and Engineering, 2022, 47(12): 15915-15930. DOI: 10.1007/s13369-022-06794-6.", 2022, True),
        ("ZIVKOVIC N, VIDAKOVIC J, MITROVIC S, et al. Implementation of dual quaternion-based robot forward kinematics algorithm in ROS[C]//2022 11th Mediterranean Conference on Embedded Computing. Piscataway: IEEE, 2022: 1-4. DOI: 10.1109/MECO55406.2022.9797160.", 2022, True),
        ("ZIVKOVIC N L J, VIDAKOVIC J Z, LAZAREVIC M P. Forward kinematics algorithm in dual quaternion space based on Denavit-Hartenberg convention[J]. Applied Engineering Letters, 2023, 8(2): 52-59. DOI: 10.18485/aeletters.2023.8.2.2.", 2023, True),
        ("MACENSKI S, FOOTE T, GERKEY B, et al. Robot Operating System 2: Design, architecture, and uses in the wild[J]. Science Robotics, 2022, 7(66): eabm6074. DOI: 10.1126/scirobotics.abm6074.", 2022, True),
        ("张明松, 黄滔. 轻型仿生机器手的几何求逆优化及轨迹规划[J]. 机械, 2022, 49(6): 73-80.", 2022, False),
        ("沈业全, 刘霞. 基于改进RRT算法的七自由度机械臂路径规划[J]. 江汉大学学报(自然科学版), 2022, 50(6): 42-49. DOI: 10.16389/j.cnki.cn42-1737/n.2022.06.005.", 2022, False),
        ("FUNDA J, TAYLOR R H, PAUL R P. On homogeneous transforms, quaternions, and computational efficiency[J]. IEEE Transactions on Robotics and Automation, 1990, 6(3): 382-388. DOI: 10.1109/70.56658.", 1990, True),
        ("KOENIG N, HOWARD A. Design and use paradigms for Gazebo, an open-source multi-robot simulator[C]//2004 IEEE/RSJ International Conference on Intelligent Robots and Systems. Piscataway: IEEE, 2004: 2149-2154. DOI: 10.1109/IROS.2004.1389727.", 2004, True),
        ("HAVILAND J, CORKE P. A systematic approach to computing the manipulator Jacobian and Hessian using the elementary transform sequence[EB/OL]. arXiv:2010.08696, 2020. DOI: 10.48550/arXiv.2010.08696.", 2020, True),
        ("DEMBY'S J, UHLMANN J, DESOUZA G N. Choosing the correct generalized inverse for the numerical solution of the inverse kinematics of incommensurate robotic manipulators[EB/OL]. arXiv:2308.02954, 2023. DOI: 10.48550/arXiv.2308.02954.", 2023, True),
        ("GUPTASARMA S, STRONG M, ZHEN H, KENNEDY M. J-PARSE: Jacobian-based projection algorithm for resolving singularities effectively in inverse kinematic control of serial manipulators[EB/OL]. arXiv:2505.00306, 2025. DOI: 10.48550/arXiv.2505.00306.", 2025, True),
    ]
    add_heading(doc, "参考文献", 1)
    for index, (text, _, _) in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-18)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(f"[{index}] {text}")
        set_run_font(run, size=9)

    recent_cutoff = 2022
    recent_count = sum(1 for _, year, _ in references if year >= recent_cutoff)
    foreign_count = sum(1 for _, _, foreign in references if foreign)
    if len(references) < 20:
        raise RuntimeError("Reference count is below 20")
    if recent_count < ceil(len(references) / 3):
        raise RuntimeError("Recent-reference ratio is below one third")
    if foreign_count < 5:
        raise RuntimeError("Foreign-language reference count is below 5")
    return len(references), recent_count, foreign_count


def apply_global_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.color.rgb = BLACK
        style.font.bold = True


def main():
    if len(TITLE_CN) > 20:
        raise RuntimeError(f"Chinese title exceeds 20 characters: {len(TITLE_CN)}")
    if len(ABSTRACT_CN) > 300:
        raise RuntimeError(f"Chinese abstract exceeds 300 characters: {len(ABSTRACT_CN)}")
    if not 3 <= len(KEYWORDS_CN) <= 8:
        raise RuntimeError("Chinese keyword count must be between 3 and 8")
    if len(KEYWORDS_CN) != len(KEYWORDS_EN):
        raise RuntimeError("Chinese and English keyword counts differ")

    doc = Document()
    apply_global_styles(doc)
    add_cover(doc)
    configure_body_section(doc)
    add_article_front(doc)
    add_main_text(doc)
    reference_count, recent_count, foreign_count = add_references(doc)

    properties = doc.core_properties
    properties.title = TITLE_CN
    properties.subject = f"《{COURSE_NAME}》研究报告"
    properties.author = AUTHOR
    properties.keywords = "; ".join(KEYWORDS_CN)
    properties.comments = "三种逆运动学算法对比、ROS 2 Jazzy与Gazebo仿真验证"

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"title_characters: {len(TITLE_CN)}")
    print(f"abstract_characters: {len(ABSTRACT_CN)}")
    print(f"keywords: {len(KEYWORDS_CN)}")
    print(f"references: {reference_count}")
    print(f"recent_references_2022_plus: {recent_count}")
    print(f"foreign_references: {foreign_count}")


if __name__ == "__main__":
    main()
