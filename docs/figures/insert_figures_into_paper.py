from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "robot_arm_kinematics_academic_dq_paper.docx"
OUTPUT = ROOT / "robot_arm_kinematics_academic_dq_paper_with_figures.docx"
FIGURES = [
    ROOT / "figures" / "fig1_system_architecture.png",
    ROOT / "figures" / "fig2_kinematics_flow.png",
    ROOT / "figures" / "fig3_validation_loop.png",
]


def replace_placeholder_with_figures():
    doc = Document(INPUT)
    figure_index = 0

    for paragraph in doc.paragraphs:
        if "[图示待生成]" not in paragraph.text:
            continue
        if figure_index >= len(FIGURES):
            raise RuntimeError("More figure placeholders than available figures")

        for run in paragraph.runs:
            run.text = ""
        paragraph.alignment = 1
        paragraph.add_run().add_picture(str(FIGURES[figure_index]), width=Inches(6.2))
        figure_index += 1

    if figure_index != len(FIGURES):
        raise RuntimeError(f"Inserted {figure_index} figures, expected {len(FIGURES)}")

    doc.save(OUTPUT)


if __name__ == "__main__":
    replace_placeholder_with_figures()
